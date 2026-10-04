# -*- coding: utf-8 -*-
"""Build competition idea deck (≤10 slides) from the official template."""
from __future__ import annotations

import io
import re
import shutil
import subprocess
import sys
import zipfile
from copy import deepcopy
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "deck" / "official" / "template.pptx"
OUT_DIR = ROOT / "deck" / "pitch"
SHOTS = OUT_DIR / "shots"
PPTX_OUT = OUT_DIR / "idea.pptx"
PDF_OUT = OUT_DIR / "idea.pdf"
NOTES_OUT = OUT_DIR / "speaker_notes.md"

# High-contrast palette for dark slides / tables
NAVY = "1A2456"
OFF_WHITE = "F2F4FF"
LIGHT_BORDER = "C8D0E8"
# Unicode bidi helpers for Latin tokens inside Arabic runs
LRM = "\u200e"
RLM = "\u200f"

# Template slide indices (0-based) to KEEP, in final order:
# 8 title, 11 problem, 13 solution, 24 pipeline, 20 reliability,
# 14 specialist, 17 results, 27 value, 18 plan, 31 closing
KEEP_ORDER = [7, 10, 12, 23, 19, 13, 16, 26, 17, 30]

NSMAP = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}


def set_shape_text(
    shape,
    text: str,
    *,
    keep_size: bool = True,
    color_hex: str | None = None,
    size_pt: float | None = None,
    bold: bool | None = None,
    rtl: bool = True,
    align=None,
) -> None:
    """Replace all paragraphs with a single run of text, preserving first-run font size."""
    tf = shape.text_frame
    first_size = None
    first_bold = None
    first_name = None
    for p in tf.paragraphs:
        for r in p.runs:
            if first_size is None and r.font.size is not None:
                first_size = r.font.size
            if first_bold is None and r.font.bold is not None:
                first_bold = r.font.bold
            if first_name is None and r.font.name:
                first_name = r.font.name
            break
        if first_size is not None:
            break
    # Clear
    tf.clear()
    p = tf.paragraphs[0]
    if rtl:
        set_paragraph_rtl(p, True)
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    elif keep_size and first_size is not None:
        run.font.size = first_size
    use_bold = bold if bold is not None else first_bold
    if use_bold is not None:
        run.font.bold = use_bold
    if first_name:
        run.font.name = first_name
    if color_hex:
        try:
            run.font.color.rgb = RGBColor.from_string(color_hex)
        except Exception:
            pass
        _force_run_srgb(run, color_hex)


def set_runs_paragraphs(shape, lines: list[str]) -> None:
    """Overwrite existing paragraphs in order; clear any extras."""
    tf = shape.text_frame
    styles = []
    for p in tf.paragraphs:
        size = bold = name = None
        for r in p.runs:
            size = r.font.size if r.font.size is not None else size
            bold = r.font.bold if r.font.bold is not None else bold
            name = r.font.name or name
        styles.append((size, bold, name))
    while len(list(tf.paragraphs)) < len(lines):
        tf.add_paragraph()
    paras = list(tf.paragraphs)
    for i, line in enumerate(lines):
        p = paras[i]
        style = styles[i] if i < len(styles) else (styles[-1] if styles else (None, None, None))
        size, bold, name = style
        if p.runs:
            p.runs[0].text = line
            for r in p.runs[1:]:
                r.text = ""
            run = p.runs[0]
        else:
            run = p.add_run()
            run.text = line
        if size is not None:
            run.font.size = size
        if bold is not None:
            run.font.bold = bold
        if name:
            run.font.name = name
    for j in range(len(lines), len(paras)):
        for r in paras[j].runs:
            r.text = ""


def find_placeholder_by_text(slide, needle: str):
    for shape in slide.shapes:
        if shape.has_text_frame and needle in (shape.text_frame.text or ""):
            return shape
    return None


def find_all_text_shapes(slide):
    return [s for s in slide.shapes if s.has_text_frame]


def set_paragraph_rtl(paragraph, rtl: bool = True) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1" if rtl else "0")


def latin(token: str) -> str:
    """Latin token with spaces + LRM (invisible; keeps LTR). Do NOT use LRI/PDI — PPT draws them as boxes."""
    return f" {LRM}{token}{LRM} "


def paren_latin(token: str) -> str:
    """'(Latin)' with LRM — keep token at end of Arabic clause without bidi glue."""
    return f"({LRM}{token}{LRM})"


def arabic_num(digits: str) -> str:
    """Digit run with LRM so ٠٤ does not flip to ٤٠ in PDF."""
    return f"{LRM}{digits}{LRM}"


def disable_autofit(shape) -> None:
    """Force noAutofit so PPT cannot shrink our sizes on export."""
    try:
        bodyPr = shape.text_frame._txBody.find(qn("a:bodyPr"))
    except Exception:
        return
    if bodyPr is None:
        return
    for child in list(bodyPr):
        tag = child.tag.split("}")[-1]
        if tag in {"normAutofit", "spAutoFit", "noAutofit"}:
            bodyPr.remove(child)
    etree.SubElement(bodyPr, qn("a:noAutofit"))


def set_mixed_text(
    shape,
    parts: list[tuple[str, bool]],
    *,
    color_hex: str | None = OFF_WHITE,
    size_pt: float | None = None,
    bold: bool | None = None,
    align=PP_ALIGN.RIGHT,
) -> None:
    """Write Arabic/Latin as separate runs so PDF keeps spaces and order."""
    tf = shape.text_frame
    # capture prior size
    first_size = None
    for p0 in tf.paragraphs:
        for r0 in p0.runs:
            if r0.font.size is not None:
                first_size = r0.font.size
                break
        if first_size is not None:
            break
    tf.clear()
    p = tf.paragraphs[0]
    set_paragraph_rtl(p, True)
    if align is not None:
        p.alignment = align
    for text, is_latin in parts:
        run = p.add_run()
        run.text = f" {text} " if is_latin else text
        if size_pt is not None:
            run.font.size = Pt(size_pt)
        elif first_size is not None:
            run.font.size = first_size
        if bold is not None:
            run.font.bold = bold
        if color_hex:
            try:
                run.font.color.rgb = RGBColor.from_string(color_hex)
            except Exception:
                pass
            _force_run_srgb(run, color_hex)


def fit_rect(box_w: int, box_h: int, img_w: int, img_h: int) -> tuple[int, int]:
    """Largest size that fits in box keeping aspect ratio."""
    scale = min(box_w / img_w, box_h / img_h)
    return int(img_w * scale), int(img_h * scale)


def set_picture_border(pic_shape, color_hex: str = LIGHT_BORDER, width_pt: float = 2.0) -> None:
    """Add a thin solid line around a picture shape."""
    spPr = pic_shape._element.spPr
    # Remove existing ln
    for ln in list(spPr.findall(qn("a:ln"))):
        spPr.remove(ln)
    ln = etree.SubElement(spPr, qn("a:ln"))
    ln.set("w", str(int(width_pt * 12700)))  # EMUs per pt
    ln.set("cap", "flat")
    ln.set("cmpd", "sng")
    solid = etree.SubElement(ln, qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", color_hex)
    etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")


def place_screenshot(slide, image_path: Path, *, box_left, box_top, box_width, box_height) -> bool:
    """Remove picture placeholder(s), place image fitted in box with light border."""
    # Drop existing picture placeholders
    for shape in list(slide.shapes):
        try:
            if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
                shape._element.getparent().remove(shape._element)
        except Exception:
            continue
    if not image_path.exists():
        print(f"WARN: missing screenshot {image_path}", file=sys.stderr)
        return False
    with Image.open(image_path) as im:
        iw, ih = im.size
    tw, th = fit_rect(int(box_width), int(box_height), iw, ih)
    left = int(box_left) + (int(box_width) - tw) // 2
    top = int(box_top) + (int(box_height) - th) // 2
    pic = slide.shapes.add_picture(str(image_path), left, top, width=tw, height=th)
    set_picture_border(pic)
    print(f"  placed screenshot {image_path.name} at {tw}×{th} emu")
    return True


def replace_picture_placeholder(slide, image_path: Path, *, enlarge: bool = False) -> bool:
    """Insert picture into the first picture placeholder (or replace its frame)."""
    for shape in slide.shapes:
        if not shape.is_placeholder:
            continue
        try:
            ph_type = shape.placeholder_format.type
        except Exception:
            continue
        if ph_type != PP_PLACEHOLDER.PICTURE:
            continue
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        if enlarge:
            # Caller should use place_screenshot; fall through to fitted replace
            shape._element.getparent().remove(shape._element)
            return place_screenshot(
                slide,
                image_path,
                box_left=left,
                box_top=top,
                box_width=width,
                box_height=height,
            )
        try:
            shape.insert_picture(str(image_path))
            print(f"  inserted picture via placeholder → {image_path.name}")
            return True
        except Exception as exc:
            print(f"  insert_picture failed ({exc}); replacing shape")
            el = shape._element
            el.getparent().remove(el)
            slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
            return True
    return False


def _force_run_srgb(run, hex_color: str) -> None:
    """Force explicit sRGB on a run; solidFill must be first child for PPT PDF export."""
    rPr = run._r.get_or_add_rPr()
    for child in list(rPr):
        if child.tag.endswith("}solidFill") or child.tag.endswith("}schemeClr"):
            rPr.remove(child)
    solid = etree.Element(qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", hex_color)
    rPr.insert(0, solid)


def set_cell_style(
    cell,
    text: str,
    *,
    fill_hex: str,
    font_hex: str,
    size_pt: float = 14,
    bold: bool = False,
    align=PP_ALIGN.CENTER,
) -> None:
    cell.text = text
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in list(tcPr.findall(qn("a:solidFill"))):
        tcPr.remove(old)
    # Kill theme/table style fills that fight our colors
    for old in list(tcPr):
        tag = old.tag
        if tag.endswith("}solidFill") or tag.endswith("}gradFill") or tag.endswith("}noFill"):
            tcPr.remove(old)
    solid = etree.SubElement(tcPr, qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", fill_hex)
    tcPr.set("anchor", "ctr")
    tf = cell.text_frame
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.alignment = align
        set_paragraph_rtl(p, True)
        for r in p.runs:
            r.font.size = Pt(size_pt)
            r.font.bold = bold
            r.font.name = "Arial"
            try:
                r.font.color.rgb = RGBColor.from_string(font_hex)
            except Exception:
                pass
            _force_run_srgb(r, font_hex)


def ensure_table_size(table, n_rows: int, n_cols: int, total_width: int | None = None) -> None:
    """Grow table rows/cols by cloning XML; force col widths to sum to total_width."""
    tbl = table._tbl
    while len(table.rows) < n_rows:
        tbl.append(deepcopy(tbl.tr_lst[-1]))
    grid = tbl.tblGrid
    while len(table.columns) < n_cols:
        grid.append(deepcopy(grid.gridCol_lst[-1]))
        for tr in tbl.tr_lst:
            tr.append(deepcopy(tr.tc_lst[-1]))
    while len(table.columns) > n_cols:
        grid.remove(grid.gridCol_lst[-1])
        for tr in tbl.tr_lst:
            tr.remove(tr.tc_lst[-1])
    while len(table.rows) > n_rows:
        tbl.remove(tbl.tr_lst[-1])
    if total_width is None:
        total_width = sum(int(gc.get("w")) for gc in grid.gridCol_lst)
    if n_cols == 5:
        weights = [2.0, 2.6, 1.1, 1.3, 1.3]
    else:
        weights = [1.0] * n_cols
    wsum = sum(weights)
    widths = [max(700_000, int(total_width * w / wsum)) for w in weights]
    widths[-1] = max(700_000, total_width - sum(widths[:-1]))
    for gc, w in zip(grid.gridCol_lst, widths):
        gc.set("w", str(w))


def demote_placeholder(shape) -> None:
    """Remove p:ph so theme placeholder colors cannot override run sRGB."""
    try:
        nvSpPr = shape._element.nvSpPr
    except AttributeError:
        return
    nvPr = nvSpPr.nvPr
    ph = nvPr.find(qn("p:ph"))
    if ph is not None:
        nvPr.remove(ph)


def add_notes(slide, text: str) -> None:
    notes = slide.notes_slide
    notes.notes_text_frame.text = text


def delete_slide(prs: Presentation, index: int) -> None:
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst[index]
    rId = sldId.get(qn("r:id"))
    prs.part.drop_rel(rId)
    sldIdLst.remove(sldId)


def reorder_slides(prs: Presentation, new_order: list[int]) -> None:
    """Reorder slides so that current indices in new_order become 0..n-1."""
    sldIdLst = prs.slides._sldIdLst
    items = list(sldIdLst)
    # Clear list
    for child in list(sldIdLst):
        sldIdLst.remove(child)
    for i in new_order:
        sldIdLst.append(items[i])


def fill_title(slide) -> None:
    # Hero image sits on the left; keep text fully clear of image + glow.
    hero = slide.shapes[1]
    hero.width = Emu(7_200_000)  # right edge ≈ 7.4M

    title = find_placeholder_by_text(slide, "[اسم المشروع]") or slide.shapes[2]
    title.left = Emu(11_000_000)
    title.top = Emu(3_100_000)
    title.width = Emu(6_400_000)
    title.height = Emu(1_800_000)
    set_shape_text(
        title,
        "فهرس مناهج التفسير",
        color_hex=OFF_WHITE,
        bold=True,
        size_pt=36,
        align=PP_ALIGN.RIGHT,
    )

    sub = find_placeholder_by_text(slide, "[وصف الفكرة") or slide.shapes[3]
    sub.left = Emu(11_000_000)
    sub.top = Emu(5_100_000)
    sub.width = Emu(6_400_000)
    sub.height = Emu(1_700_000)
    set_shape_text(
        sub,
        "كيف فسّر المفسِّر هذه الآية؟ الجواب من نصه، بالدليل",
        color_hex=OFF_WHITE,
        size_pt=16,
        align=PP_ALIGN.RIGHT,
    )

    # Plain RTL track line — digits as separate LTR-ish run so ٠٤ stays correct
    for shape in slide.shapes:
        t = shape.text_frame.text if shape.has_text_frame else ""
        if "[اسم الفريق" in t or "اسم الفريق" in t or t.startswith("المسار"):
            # Keep clear of bottom-right cyan glow; sit on navy
            shape.left = Emu(8_800_000)
            shape.top = Emu(8_750_000)
            shape.width = Emu(6_800_000)
            shape.height = Emu(450_000)
            set_mixed_text(
                shape,
                [("المسار ", False), ("٠٤", True), (" — أدوات المعرفة والتحقق", False)],
                color_hex=OFF_WHITE,
                size_pt=13,
                align=PP_ALIGN.RIGHT,
            )
        if "[تاريخ" in t or t.startswith("سبتمبر"):
            set_shape_text(shape, "سبتمبر 2026", color_hex=OFF_WHITE, size_pt=12)
    add_notes(
        slide,
        "نفتح باسم المنتج: فهرس مناهج التفسير. سؤال واحد: كيف فسّر المفسِّر هذه الآية؟ "
        "الجواب من نصه وبالدليل. المسار ٠٤: أدوات المعرفة والتحقق. لا أسماء أشخاص.",
    )


def fill_problem(slide) -> None:
    shapes = {s.text_frame.text.strip()[:20]: s for s in find_all_text_shapes(slide) if s.has_text_frame}
    # Map by known placeholders
    mapping = [
        ("[اسم القسم]", "المشكلة"),
        ("[عنوان الشريحة]", "المعرفة موجودة… والوصول إليها بطيء وغير موثّق"),
        ("[فكرة واحدة", "الباحث يقرأ صفحات ليعرف منهج المفسّر؛ وإجابات الذكاء عن القرآن بلا أثر قابل للتتبع"),
        ("[النقطة الأولى]", "قراءة طويلة"),
        ("[شرح قصير يدعم الفكرة مع دليل أو مثال]", "ليُعرف أين استُخدم القرآن أو السنة أو أقوال الصحابة في التفسير"),
        ("[النقطة الثانية]", "ذكاء بلا إسناد"),
        ("[شرح قصير يدعم الفكرة مع دليل أو مثال]", "إجابات عن القرآن بلا مصدر ظاهر ولا حالة دليل"),
        ("[النقطة الثالثة]", "معيار المسار 04"),
        ("[شرح قصير يدعم الفكرة مع دليل أو مثال]", "«هل حسّن الحل دقة الوصول إلى المعرفة أو التحقق منها»"),
        ("[النقطة الرابعة]", "ما نحتاجه"),
        ("[شرح قصير يدعم الفكرة مع دليل أو مثال]", "مصدر ظاهر، حالة دليل، وتمييز المؤيَّد عما يحتاج تحققاً أو إحالة"),
    ]
    # More robust: iterate shapes in order of known indices from inspection
    # slide 11 shape indices:
    # 1 section, 2 title, 3 idea, 5 p1, 6 d1, 7 p2, 8 d2, 9 p3, 10 d3, 11 p4, 12 d4
    texts = [
        (1, "المشكلة"),
        (2, "المعرفة موجودة… والوصول بطيء وغير موثّق"),
        (3, "الباحث يقرأ صفحات ليعرف منهج المفسّر؛ وإجابات الذكاء عن القرآن بلا أثر"),
        (5, "قراءة طويلة"),
        (6, "ليُعرف أين استُخدم القرآن أو السنة أو الصحابة في النص"),
        (7, "ذكاء بلا إسناد"),
        (8, "إجابات عن القرآن بلا مصدر ظاهر ولا حالة دليل"),
        (9, "معيار نجاح المسار"),
        (10, "«هل حسّن الحل دقة الوصول إلى المعرفة أو التحقق منها»"),
        (11, "ما يُطلب"),
        (12, "مصدر ظاهر، حالة دليل، وتمييز المؤيَّد عما يحتاج تحققاً أو إحالة"),
    ]
    for idx, text in texts:
        set_shape_text(slide.shapes[idx], text)
    add_notes(
        slide,
        "المشكلة مزدوجة: الباحث يغرق في الصفحات ليلتقط منهج المفسّر، وأدوات الذكاء تجيب عن القرآن "
        "بلا أثر يُراجع. معيار المسار: هل حسّن الحل دقة الوصول أو التحقق؟ نريد مصدراً ظاهراً وحالة دليل.",
    )


def fill_solution(slide) -> None:
    # slide 13: 3 section, 4 title, 7 h1, 8 d1, 9 h2, 10 d2; pic placeholder 6
    set_shape_text(slide.shapes[3], "الحل")
    set_shape_text(slide.shapes[4], "الآلة تقترح، والمتخصص يحكم، والنص لا يتغيّر حرفاً")
    set_shape_text(slide.shapes[7], "ما الذي نراه؟")
    set_shape_text(
        slide.shapes[8],
        "صفحة التفسير يميناً ولوحة المناهج يساراً؛ تمييز ملوّن بدليله",
    )
    set_shape_text(slide.shapes[9], "ما الذي يهم؟")
    set_shape_text(
        slide.shapes[10],
        "اقتراح آلي للمراجعة؛ الاعتماد للمختص؛ حرف المصدر ثابت",
    )
    # Enlarge screenshot into left column under title; keep aspect; light border
    ok = place_screenshot(
        slide,
        SHOTS / "fahras_ibn_kathir.jpg",
        box_left=1_050_000,
        box_top=2_750_000,
        box_width=9_100_000,
        box_height=6_100_000,
    )
    if not ok:
        print("WARN: could not insert solution screenshot", file=sys.stderr)
    add_notes(
        slide,
        "هذه صورة المنتج: نص المفسّر ملوناً بمناهجه، والآلة تقترح والمتخصص يحكم. "
        "الجملة الحاكمة: الآلة تقترح، والمتخصص يحكم، والنص لا يتغيّر حرفاً.",
    )


def fill_pipeline(slide) -> None:
    set_shape_text(slide.shapes[1], "آلية العمل", color_hex=OFF_WHITE)
    # Latin token at end in parentheses — avoids «JSONمعتمد» bidi glue
    set_mixed_text(
        slide.shapes[2],
        [
            ("من النص المُثبَّت إلى ملف معتمد (", False),
            ("JSON", True),
            (")", False),
        ],
        color_hex=OFF_WHITE,
        align=PP_ALIGN.RIGHT,
    )
    set_shape_text(
        slide.shapes[3],
        "النموذج قابل للاستبدال؛ الضمانات ثابتة — لم ندرّب نموذجاً بعد؛ بنينا منهجاً (سكيل)",
        color_hex=OFF_WHITE,
    )
    # cards 5..8
    cards = [
        (5, ["01", "المدخلات", "نص ببصمة رقمية", "أجزاء مرقّمة + فصل حاشية المحقق"]),
        (6, ["02", "المعالجة", "علامات نصية بلا ذكاء", "ثم النموذج يقترح بأرقام الأجزاء فقط"]),
        (7, ["03", "المخرجات", "درجة المدقق المستقل", "توجيه للمختص أو مرشّح آلي"]),
        (8, ["04", "الأثر", "قلم التمييز + اعتماد", f"ملف معتمد {paren_latin('JSON')}؛ نموذج ضعيف → مزيد للمختص لا فساد"]),
    ]
    for idx, lines in cards:
        set_runs_paragraphs(slide.shapes[idx], lines)
    add_notes(
        slide,
        "السلسلة: نص ببصمة، أجزاء وفصل الحاشية، علامات بلا ذكاء، اقتراح بأرقام الأجزاء، "
        "مدقق مستقل، ثم قلم المتخصص. أي نموذج API يندرج؛ الضمانات لا تتغير. لم ندرّب نموذجاً بعد.",
    )


def fill_reliability(slide) -> None:
    # layout 20 light 4-points: shapes 1..12 similar to 11
    set_shape_text(slide.shapes[1], "الموثوقية")
    set_shape_text(slide.shapes[2], "الموثوقية والسلامة العلمية")
    set_shape_text(
        slide.shapes[3],
        "نفهرس منهج المفسّر لا الفتوى؛ الحساس للمختص؛ امتناع عند غياب الدليل",
    )
    set_shape_text(slide.shapes[5], "مستويات المحتوى (أ–د)")
    set_shape_text(
        slide.shapes[6],
        "أ معلومات مستقرة · ب شرح · ج حساسة → مختص · د فتوى → لا حكم مستقل",
    )
    set_shape_text(slide.shapes[7], "الوظيفة لا الألفاظ")
    set_shape_text(
        slide.shapes[8],
        "آية داخل حديث ≠ قرآن بالقرآن؛ صحابي في إسناد ≠ قول صحابي؛ يقين: صريح/مستنتج/امتناع",
    )
    set_shape_text(slide.shapes[9], "مقترح آلي ≠ معتمد")
    set_shape_text(
        slide.shapes[10],
        "التمييز مقترح حتى يعتمد المتخصص؛ «سنة» استعمال حديث لا حكم صحة",
    )
    set_shape_text(slide.shapes[11], "المصدر")
    # Same size class as the other three body blocks (template ≈ 24 pt)
    src = slide.shapes[12]
    src.height = Emu(1_400_000)  # room for 24 pt wrap
    set_mixed_text(
        src,
        [
            # Latin only inside trailing parens — avoids «4.0 CC BY» reorder
            ("نص مركز تفسير — بيانات مفتوحة (", False),
            ("CC BY 4.0", True),
            (") · نبدأ بالطبري (ت ", False),
            ("٣١٠", True),
            ("هـ) · الأداة تدرس منهج المفسِّر ولا تتخذ التفسير مصدراً لإجابة أو فتوى", False),
        ],
        color_hex=None,
        size_pt=24,
        align=PP_ALIGN.RIGHT,
    )
    disable_autofit(src)
    add_notes(
        slide,
        "نربط الوسم بشاهد في نص مثبت، ونفصل كلام المفسّر عن الحاشية وعن مقترح الذكاء. "
        "المستويات الأربعة للحزمة: أداتنا في فهرسة المنهج؛ الجيم للمختص؛ الدال بلا حكم. "
        "نمتنع عند غياب الدليل. لا ندّعي صحة حديث بالوسم.",
    )


def fill_specialist(slide) -> None:
    # slide 14: 3 section, 4 title, 7 idea, 8 body, 9 caption; pic 6
    set_shape_text(slide.shapes[3], "أداة المتخصص")
    set_shape_text(slide.shapes[4], "قلم تمييز ذكي + مقارنة بالأصل")
    set_shape_text(slide.shapes[7], "تمديد وحدف وإعادة تلوين كلمةً كلمة")
    set_shape_text(
        slide.shapes[8],
        "قارن بالمصدر قبل الاعتماد؛ كل تعديل يُسجَّل → دقة لكل منهج وبيانات تدريب معتمدة لاحقاً",
    )
    set_shape_text(slide.shapes[9], "لقطة: الواجهة + لوحة المناهج")
    img = SHOTS / "methods.jpg"
    if not img.exists():
        img = SHOTS / "fahras_ibn_kathir.jpg"
    ok = place_screenshot(
        slide,
        img,
        box_left=1_050_000,
        box_top=2_600_000,
        box_width=8_300_000,
        box_height=6_200_000,
    )
    if not ok:
        print("WARN: could not insert specialist screenshot", file=sys.stderr)
    add_notes(
        slide,
        "المتخصص يصحّح الحدود والمنهج دون تغيير حرف المصدر، ويقارن بالأصل قبل الاعتماد. "
        "السجل يقيس الدقة لكل منهج ويجهّز بيانات تدريب لاحقة داخل الضمانات نفسها.",
    )


def fill_results(slide) -> None:
    # Dark table → high-contrast 5-col results
    set_shape_text(slide.shapes[2], "النتائج", color_hex=OFF_WHITE)
    set_shape_text(
        slide.shapes[3],
        "تشغيل واحد، أربعة تفاسير، قواعد موحّدة",
        color_hex=OFF_WHITE,
    )
    set_mixed_text(
        slide.shapes[4],
        [
            ("مدقق صارم؛ بلا ضبط لكل تفسير · راجعنا أنفسنا: متساهل فشدّدنا ونشرنا الاثنين (", False),
            ("Run 1", True),
            (" → ", True),
            ("Run 2", True),
            (")", False),
        ],
        color_hex=OFF_WHITE,
        size_pt=14,
        align=PP_ALIGN.RIGHT,
    )
    table_shape = None
    for shape in slide.shapes:
        if shape.has_table:
            table_shape = shape
            break
    assert table_shape is not None
    table = table_shape.table

    # Fit inside slide; grid widths MUST equal shape width (else last col clips)
    table_w = 15_200_000
    table_shape.left = Emu(1_400_000)
    table_shape.top = Emu(3_650_000)
    table_shape.width = Emu(table_w)
    table_shape.height = Emu(4_400_000)

    ensure_table_size(table, 6, 5, total_width=table_w)

    # Drop template table style so our fills/colors stick in PDF
    tblPr = table._tbl.tblPr
    if tblPr is not None:
        for child in list(tblPr):
            if child.tag.endswith("}tableStyleId"):
                tblPr.remove(child)

    headers = ["التفسير", "النموذج", "مواضع", "قبول آلي", "للمتخصص"]
    rows = [
        headers,
        ["ابن كثير", f"النموذج أ {paren_latin('Codex')}", "42", "23", "19"],
        ["ابن كثير", f"النموذج ب {paren_latin('DeepSeek')}", "44", "26", "18"],
        # Latin tokens only inside trailing parens — no Arabic glued after ')'
        ["الطبري", f"جزئياً {paren_latin('DeepSeek + MiMo')}", "31", "18", "13"],
        ["البغوي", paren_latin("DeepSeek"), "39", "29", "10"],
        ["السعدي", paren_latin("DeepSeek"), "24", "15", "9"],
    ]
    for r_i, row_vals in enumerate(rows):
        is_header = r_i == 0
        for c_i, val in enumerate(row_vals):
            cell = table.cell(r_i, c_i)
            numeric = (not is_header) and c_i >= 2
            if is_header:
                # Off-white header + navy text: survives PPT theme overrides that keep dark glyphs
                set_cell_style(
                    cell,
                    val,
                    fill_hex=OFF_WHITE,
                    font_hex=NAVY,
                    size_pt=16,
                    bold=True,
                )
            elif numeric:
                set_cell_style(
                    cell,
                    val,
                    fill_hex=OFF_WHITE,
                    font_hex=NAVY,
                    size_pt=22,
                    bold=True,
                )
            else:
                # التفسير / النموذج: ≥ 16 pt, bold on name col (same weight class as numbers)
                set_cell_style(
                    cell,
                    val,
                    fill_hex=OFF_WHITE,
                    font_hex=NAVY,
                    size_pt=18,
                    bold=True,
                    align=PP_ALIGN.RIGHT if c_i < 2 else PP_ALIGN.CENTER,
                )

    footnote = slide.shapes[7]
    footnote.top = Emu(8_250_000)
    footnote.height = Emu(800_000)
    set_shape_text(
        footnote,
        "١٠٠٪ تطابق حرفي مع المصدر المُثبَّت · ٢٬٠٠٦ مقطعاً موثّقاً · نموذجان متقاربان على ابن كثير · "
        "أرقام تجربة أولى على ٣ آيات؛ الدقة تُقاس بعيّنة المتخصص",
        color_hex=OFF_WHITE,
        size_pt=12,
    )
    add_notes(
        slide,
        "أرقام Run 2 كما هي: ابن كثير ٤٢/٢٣/١٩ و٤٤/٢٦/١٨، الطبري ٣١/١٨/١٣، البغوي ٣٩/٢٩/١٠، "
        "السعدي ٢٤/١٥/٩. راجعنا أنفسنا بعد Run 1 المتساهل. لا نثق برأي النموذج في نفسه. "
        "النموذجان على ابن كثير متقاربان.",
    )


def fill_value(slide) -> None:
    # slide 27 elements — shapes 2 section, 3 title, 4 subtitle, then pairs
    set_shape_text(slide.shapes[2], "الأصالة والقيمة")
    set_shape_text(slide.shapes[3], "كيف فسّر… لا ماذا قيل فقط")
    set_shape_text(
        slide.shapes[4],
        "المواقع الكبرى تعرض ماذا قيل عن الآية؛ نحن نكشف كيف فسّرها المفسّر في نصه، بالدليل",
    )
    # Correct indices from inspection: 7/8, 10/11, 13/14, 16/17
    # JSON as separate LTR run inside trailing parens (avoids «)دون» glue)
    pairs = [
        (
            7,
            8,
            "للباحث",
            [("تصفية بالمنهج عبر التفاسير؛ لا نعلم نظاماً منشوراً يُعلّم الموضع بالدليل هكذا", False)],
        ),
        (
            10,
            11,
            "للمعلّم والقارئ",
            [("بطاقة قصيرة معتمدة؛ القارئ العادي يرى المعتمد فقط", False)],
        ),
        (
            13,
            14,
            "للحفظ والتطبيقات",
            [
                ("كل آية محفوظة مربوطة بشرح قرآن بالقرآن المعتمد؛ ملف معتمد (", False),
                ("JSON", True),
                (") دون ذكاء وقت القراءة", False),
            ],
        ),
        (
            16,
            17,
            "امتداد المتصفح",
            [
                ("لوحة من ملف معتمد (", False),
                ("JSON", True),
                (") · تلوين إن طابقت الطبعة · اقتراح آلي غير معتمد؛ بلا تعديل لنص الموقع", False),
            ],
        ),
    ]
    for t_i, d_i, title, parts in pairs:
        set_shape_text(slide.shapes[t_i], title)
        if len(parts) == 1 and not parts[0][1]:
            set_shape_text(slide.shapes[d_i], parts[0][0], color_hex=OFF_WHITE)
        else:
            set_mixed_text(
                slide.shapes[d_i],
                parts,
                color_hex=OFF_WHITE,
                align=PP_ALIGN.RIGHT,
            )
    add_notes(
        slide,
        "الفرق: المواقع تعرض ماذا قيل؛ نحن نكشف كيف فسّر المفسّر في نصه بالدليل. "
        "قيمة للباحث والمعلّم والقارئ والحفظ وتطبيقات القرآن وامتداد المتصفح. العادي يرى المعتمد فقط.",
    )


def fill_plan(slide) -> None:
    set_shape_text(slide.shapes[2], "خطة البناء")
    set_shape_text(slide.shapes[3], "أيام البناء ٤–٦ أكتوبر وما بعدها")
    set_shape_text(
        slide.shapes[4],
        "مرور كامل على الأنفال + تدقيق ١٠٪ + مختص للحساس · مستودع عام · عرض حي · فيديو",
    )
    # 5 columns: shapes 5-8 (01), 9-12 (02), 13-16 (03), 17-20 (04), 21-24 (05)
    stages = [
        (5, "01", "اليوم ١", "٤ أكتوبر", "الأنفال: تصنيف كامل"),
        (9, "02", "اليوم ٢", "٥ أكتوبر", "تدقيق عشوائي ١٠٪ للآلي"),
        (13, "03", "اليوم ٣", "٦ أكتوبر", "مختص: إسرائيليات وأسباب نزول ومغازي"),
        (
            17,
            "04",
            "التسليم",
            "مستودع + عرض",
            # Avoid wrapping mid «Live Demo»; avoid bare ≤ (RTL often mirrors it to ≥)
            f"مستودع عام {paren_latin('GitHub')} · عرض حي · فيديو بحدّ ٢د",
        ),
        (21, "05", "بعدها", "خارطة طريق", "مرجع مُحكِّم · تفاسير مرخّصة · «لم ندرّب نموذجاً بعد» · أداة للمراكز"),
    ]
    for base, num, stage, period, out in stages:
        set_shape_text(slide.shapes[base], num)
        set_shape_text(slide.shapes[base + 1], stage)
        set_shape_text(slide.shapes[base + 2], period)
        set_shape_text(slide.shapes[base + 3], out)
    add_notes(
        slide,
        "في أيام البناء: سورة الأنفال كاملة، تدقيق عشرة بالمئة، ومختص للحساس. "
        "التسليم: مستودع عام وعرض حي وفيديو. لاحقاً المرجع المُحكِّم والمزيد من التفاسير عند الترخيص، "
        "ونموذج مفتوح على الاعتمادات — لم ندرّب بعد — وأداة مفتوحة لمراكز التفسير.",
    )


def fill_closing(slide) -> None:
    # Layout top→bottom with clear gaps: verse (hero) · tagline · credits
    # Amiri not installed on this Windows host → embed-safe Traditional Arabic
    # Traditional Arabic reads smaller than sans at same pt — 72 (≥ 54), clear of logo
    verse_font = "Traditional Arabic"
    verse_pt = 72
    verse = slide.shapes[3]
    verse.left = Emu(2_400_000)
    verse.top = Emu(3_200_000)
    verse.width = Emu(13_400_000)
    verse.height = Emu(2_000_000)
    set_shape_text(
        verse,
        "﴿وَبِالْحَقِّ أَنزَلْنَاهُ وَبِالْحَقِّ نَزَلَ﴾",
        color_hex=OFF_WHITE,
        size_pt=verse_pt,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    verse.text_frame.word_wrap = True
    verse.text_frame.auto_size = None
    try:
        verse.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    except Exception:
        pass
    for p in verse.text_frame.paragraphs:
        set_paragraph_rtl(p, True)
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.name = verse_font
            r.font.size = Pt(verse_pt)
            r.font.bold = True
            try:
                r.font.color.rgb = RGBColor.from_string(OFF_WHITE)
            except Exception:
                pass
            rPr = r._r.get_or_add_rPr()
            for tag, typeface in (
                (qn("a:latin"), verse_font),
                (qn("a:ea"), verse_font),
                (qn("a:cs"), verse_font),
            ):
                el = rPr.find(tag)
                if el is None:
                    el = etree.SubElement(rPr, tag)
                el.set("typeface", typeface)
            _force_run_srgb(r, OFF_WHITE)
    disable_autofit(verse)

    tag = slide.shapes[1]
    tag.left = Emu(2_400_000)
    tag.top = Emu(5_500_000)
    tag.width = Emu(13_400_000)
    tag.height = Emu(700_000)
    set_shape_text(
        tag,
        "نحفظ لنصّ المفسِّر الأمانة",
        color_hex=OFF_WHITE,
        size_pt=28,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    for p in tag.text_frame.paragraphs:
        for r in p.runs:
            _force_run_srgb(r, OFF_WHITE)
    disable_autofit(tag)

    cred = slide.shapes[2]
    cred.left = Emu(2_000_000)
    cred.top = Emu(6_500_000)
    cred.width = Emu(14_200_000)
    cred.height = Emu(900_000)
    # Latin tokens at end of each clause in parentheses
    set_mixed_text(
        cred,
        [
            ("مصدر النص: مركز تفسير — بيانات مفتوحة (", False),
            ("CC BY 4.0", True),
            (") · الشفرة (", False),
            ("MIT", True),
            (") · أدوات الذكاء مفصح عنها في المستودع", False),
        ],
        color_hex=OFF_WHITE,
        size_pt=14,
        align=PP_ALIGN.CENTER,
    )
    for p in cred.text_frame.paragraphs:
        for r in p.runs:
            _force_run_srgb(r, OFF_WHITE)
    disable_autofit(cred)
    # Keep as placeholders (like slide 1): demoting made PPT PDF render black text
    add_notes(
        slide,
        "نختم بالآية: وبالحق أنزلناه وبالحق نزل. نحفظ لنصّ المفسّر الأمانة. "
        "الائتمان: مركز تفسير CC BY 4.0، الشفرة MIT، وأدوات الذكاء مفصح عنها.",
    )


FILLERS = [
    fill_title,
    fill_problem,
    fill_solution,
    fill_pipeline,
    fill_reliability,
    fill_specialist,
    fill_results,
    fill_value,
    fill_plan,
    fill_closing,
]

SPEAKER_NOTES = [
    ("1 — العنوان", "نفتح باسم المنتج: فهرس مناهج التفسير. سؤال واحد: كيف فسّر المفسِّر هذه الآية؟ الجواب من نصه وبالدليل. المسار 04."),
    ("2 — المشكلة", "الباحث يغرق في الصفحات؛ والذكاء يجيب بلا أثر. معيار المسار: هل حسّن الحل دقة الوصول أو التحقق؟"),
    ("3 — الحل", "صورة المنتج. الآلة تقترح، والمتخصص يحكم، والنص لا يتغيّر حرفاً."),
    ("4 — آلية العمل", "بصمة وأجزاء وحاشية وعلامات ثم اقتراح بأرقام أجزاء فمدقق فقلم متخصص. النموذج قابل للاستبدال."),
    ("5 — الموثوقية", "مستويات أ–د، الوظيفة لا الألفاظ، مقترح مقابل معتمد، لا حكم حديث من الوسم، وامتناع عند غياب الدليل."),
    ("6 — أداة المتخصص", "قلم تمييز ومقارنة بالأصل وسجل يقيس الدقة."),
    ("7 — النتائج", "أرقام Run 2 على أربعة تفاسير، ومراجعة أنفسنا بعد Run 1، ونموذجان متقاربان على ابن كثير."),
    ("8 — الأصالة والقيمة", "كيف فسّر لا ماذا قيل فقط؛ باحث ومعلّم وقارئ وحفظ وتطبيقات وامتداد."),
    ("9 — خطة البناء", "الأنفال وتدقيق ١٠٪ ومختص للحساس؛ مستودع وعرض وفيديو؛ خارطة الطريق بلا تدريب نموذج بعد."),
    ("10 — الختام", "وبالحق أنزلناه وبالحق نزل. نحفظ الأمانة. مركز تفسير CC BY 4.0 والشفرة MIT."),
]


def compress_media_in_pptx(pptx_path: Path, max_bytes: int = 9_500_000) -> None:
    """Recompress large media and drop orphans; rewrite pptx in place."""
    tmp = pptx_path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(pptx_path, "r") as zin:
        # Collect referenced media
        referenced = set()
        for name in zin.namelist():
            if name.startswith("ppt/slides/_rels/") or name.startswith("ppt/slideLayouts/_rels/") or name.startswith("ppt/slideMasters/_rels/"):
                data = zin.read(name).decode("utf-8", errors="ignore")
                for m in re.findall(r"media/([^\"']+)", data):
                    referenced.add(m)
            if name == "ppt/_rels/presentation.xml.rels":
                data = zin.read(name).decode("utf-8", errors="ignore")
                for m in re.findall(r"media/([^\"']+)", data):
                    referenced.add(m)

        # Also scan slide XML for embeds
        for name in zin.namelist():
            if "/media/" in name:
                continue

        content_types = zin.read("[Content_Types].xml")
        ct_root = etree.fromstring(content_types)

        out_buf = io.BytesIO()
        with zipfile.ZipFile(out_buf, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("ppt/media/"):
                    media_name = item.filename.split("/")[-1]
                    if media_name not in referenced:
                        print(f"  drop unused media {media_name}")
                        continue
                    # recompress large raster
                    lower = media_name.lower()
                    if lower.endswith((".png", ".jpg", ".jpeg")) and len(data) > 120_000:
                        try:
                            im = Image.open(io.BytesIO(data))
                            im = im.convert("RGB")
                            w, h = im.size
                            # title/closing backgrounds can be wide; cap 1920 on long edge
                            max_edge = 1600 if len(data) > 500_000 else 1920
                            if max(w, h) > max_edge:
                                if w >= h:
                                    im = im.resize((max_edge, int(h * max_edge / w)), Image.Resampling.LANCZOS)
                                else:
                                    im = im.resize((int(w * max_edge / h), max_edge), Image.Resampling.LANCZOS)
                            bio = io.BytesIO()
                            # Prefer JPEG for photos/gradients
                            q = 72 if len(data) > 1_000_000 else 80
                            im.save(bio, format="JPEG", quality=q, optimize=True)
                            new_data = bio.getvalue()
                            if len(new_data) < len(data):
                                print(f"  recompress {media_name}: {len(data)} → {len(new_data)}")
                                data = new_data
                                # If still .png name but JPEG bytes, update content type override
                                if lower.endswith(".png"):
                                    # Add Override for this part as jpeg
                                    part_name = "/" + item.filename
                                    # Remove existing Override for this part
                                    for ov in list(ct_root.findall("{http://schemas.openxmlformats.org/package/2006/content-types}Override")):
                                        if ov.get("PartName") == part_name:
                                            ct_root.remove(ov)
                                    ov = etree.SubElement(
                                        ct_root,
                                        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                                    )
                                    ov.set("PartName", part_name)
                                    ov.set("ContentType", "image/jpeg")
                        except Exception as exc:
                            print(f"  skip recompress {media_name}: {exc}")
                if item.filename == "[Content_Types].xml":
                    continue  # write updated later
                zout.writestr(item, data)
            zout.writestr(
                "[Content_Types].xml",
                etree.tostring(ct_root, xml_declaration=True, encoding="UTF-8", standalone=True),
            )

    tmp.write_bytes(out_buf.getvalue())
    tmp.replace(pptx_path)
    size = pptx_path.stat().st_size
    print(f"pptx after media pass: {size} bytes")
    if size > max_bytes:
        # second harsher pass
        _harsh_recompress(pptx_path)


def _harsh_recompress(pptx_path: Path) -> None:
    tmp = pptx_path.with_suffix(".tmp2.pptx")
    with zipfile.ZipFile(pptx_path, "r") as zin, zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("ppt/media/") and len(data) > 80_000:
                try:
                    im = Image.open(io.BytesIO(data)).convert("RGB")
                    w, h = im.size
                    max_edge = 1280
                    if max(w, h) > max_edge:
                        if w >= h:
                            im = im.resize((max_edge, max(1, int(h * max_edge / w))), Image.Resampling.LANCZOS)
                        else:
                            im = im.resize((max(1, int(w * max_edge / h)), max_edge), Image.Resampling.LANCZOS)
                    bio = io.BytesIO()
                    im.save(bio, format="JPEG", quality=60, optimize=True)
                    if len(bio.getvalue()) < len(data):
                        print(f"  harsh {item.filename}: {len(data)} → {len(bio.getvalue())}")
                        data = bio.getvalue()
                except Exception:
                    pass
            zout.writestr(item, data)
    tmp.replace(pptx_path)
    print(f"pptx after harsh pass: {pptx_path.stat().st_size} bytes")


def export_pdf(pptx_path: Path, pdf_path: Path) -> bool:
    # Try LibreOffice
    soffice_candidates = [
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        shutil.which("soffice"),
    ]
    for c in soffice_candidates:
        if not c:
            continue
        c = Path(c)
        if c.exists():
            outdir = pdf_path.parent
            cmd = [str(c), "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(pptx_path)]
            print("LibreOffice:", " ".join(cmd))
            subprocess.run(cmd, check=False, timeout=180)
            # soffice names output after input stem
            produced = outdir / (pptx_path.stem + ".pdf")
            if produced.exists() and produced != pdf_path:
                produced.replace(pdf_path)
            if pdf_path.exists():
                return True

    # PowerPoint COM — force light font on closing slide before PDF (COM sometimes drops sRGB)
    try:
        ps = f"""
$ppt = New-Object -ComObject PowerPoint.Application
$ppt.Visible = [Microsoft.Office.Core.MsoTriState]::msoTrue
$pres = $ppt.Presentations.Open('{pptx_path.resolve()}', $false, $false, $false)
# BGR for #F2F4FF
$light = 0xFFF4F2
foreach ($si in 1..$pres.Slides.Count) {{
  $slide = $pres.Slides.Item($si)
  # Dark template slides that need light text
  if ($si -in 1,3,4,6,7,8,9,10) {{
    for ($i = 1; $i -le $slide.Shapes.Count; $i++) {{
      $sh = $slide.Shapes.Item($i)
      try {{
        if ($sh.HasTextFrame -eq -1 -and $sh.TextFrame.HasText -eq -1) {{
          if ($sh.HasTable -eq -1) {{ continue }}
          $tr = $sh.TextFrame.TextRange
          # Skip footer partner label teal; force light on body copy
          if ($tr.Text -notmatch 'شركاؤنا') {{
            $tr.Font.Color.RGB = $light
          }}
        }}
      }} catch {{}}
    }}
  }}
}}
$pres.SaveAs('{pdf_path.resolve()}', 32)
$pres.Close()
$ppt.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($pres) | Out-Null
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null
"""
        print("Trying PowerPoint COM…")
        r = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps],
            capture_output=True,
            text=True,
            timeout=180,
        )
        if pdf_path.exists():
            return True
        print("COM stderr:", r.stderr[-500:] if r.stderr else "")
    except Exception as exc:
        print("COM failed:", exc)
    return pdf_path.exists()


def scan_brackets(prs: Presentation) -> list[str]:
    hits = []
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text or ""
                for m in re.findall(r"\[[^\]]+\]", t):
                    hits.append(f"slide {i}: {m}")
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for m in re.findall(r"\[[^\]]+\]", cell.text or ""):
                            hits.append(f"slide {i} table: {m}")
    return hits


def scan_min_font_sizes(prs: Presentation, min_pt: float = 12.0) -> list[str]:
    """Per-slide: report any run/cell text smaller than min_pt (skip empty / page-number)."""
    issues: list[str] = []
    per_slide_min: dict[int, float] = {}
    for i, slide in enumerate(prs.slides, 1):
        smallest = None
        for shape in slide.shapes:
            frames = []
            if shape.has_text_frame:
                frames.append(shape.text_frame)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        frames.append(cell.text_frame)
            for tf in frames:
                text = (tf.text or "").strip()
                if not text or text in {"‹#›", "#"}:
                    continue
                for p in tf.paragraphs:
                    for r in p.runs:
                        if not (r.text or "").strip():
                            continue
                        if r.font.size is None:
                            continue
                        pt = r.font.size.pt
                        if smallest is None or pt < smallest:
                            smallest = pt
                        if pt + 1e-6 < min_pt:
                            snippet = (r.text or "").replace("\n", " ")[:40]
                            issues.append(f"slide {i}: {pt:g} pt < {min_pt:g} :: {snippet!r}")
        if smallest is not None:
            per_slide_min[i] = smallest
    for i in sorted(per_slide_min):
        print(f"  slide {i} min font: {per_slide_min[i]:g} pt")
    return issues


def verify_pdf_overflow(pdf_path: Path, margin_pt: float = 8.0) -> list[str]:
    """Render each page; flag text blocks that clip near/over page edges."""
    import fitz

    issues: list[str] = []
    doc = fitz.open(pdf_path)
    out_dir = OUT_DIR / "_pdf_preview"
    out_dir.mkdir(exist_ok=True)
    for page_i, page in enumerate(doc, 1):
        rect = page.rect
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        png = out_dir / f"slide_{page_i:02d}.png"
        pix.save(str(png))
        blocks = page.get_text("blocks")
        for b in blocks:
            x0, y0, x1, y1, text, *_ = b
            text = (text or "").strip()
            if not text:
                continue
            clipped = []
            if x0 < margin_pt:
                clipped.append(f"left({x0:.1f})")
            if y0 < margin_pt:
                clipped.append(f"top({y0:.1f})")
            if x1 > rect.width - margin_pt:
                clipped.append(f"right({x1:.1f}>{rect.width - margin_pt:.1f})")
            if y1 > rect.height - margin_pt:
                clipped.append(f"bottom({y1:.1f}>{rect.height - margin_pt:.1f})")
            if clipped:
                snippet = text.replace("\n", " ")[:60]
                issues.append(f"slide {page_i}: {', '.join(clipped)} :: {snippet!r}")
        # Also dump a few key strings for bidi spot-check
        raw = page.get_text("text")
        if page_i == 1 and "المسار" in raw:
            for line in raw.splitlines():
                if "المسار" in line:
                    print(f"  PDF s1 track line: {line!r}")
        if page_i in (4, 7, 8, 9, 10):
            for line in raw.splitlines():
                if any(tok in line for tok in ("JSON", "Run", "GitHub", "Live Demo", "Codex", "DeepSeek", "MIT", "CC BY", "حق", "نحفظ")):
                    print(f"  PDF s{page_i}: {line!r}")
        # Flag tiny PDF text (< 12 pt) via span sizes
        for block in page.get_text("dict").get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    size = float(span.get("size") or 0)
                    txt = (span.get("text") or "").strip()
                    if not txt or size <= 0:
                        continue
                    if size + 1e-6 < 12.0:
                        issues.append(
                            f"slide {page_i}: tiny PDF span {size:.1f}pt :: {txt.replace(chr(10), ' ')[:50]!r}"
                        )
    doc.close()
    return issues


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if not TEMPLATE.exists():
        print("Missing template", TEMPLATE)
        return 1
    for need in ("fahras_ibn_kathir.jpg", "methods.jpg"):
        if not (SHOTS / need).exists():
            print("Missing shot", need)

    print("Loading template…")
    prs = Presentation(str(TEMPLATE))
    assert len(prs.slides) == 31

    # Fill kept slides first (while indices stable)
    for order_i, slide_idx in enumerate(KEEP_ORDER):
        print(f"Filling template slide {slide_idx + 1} → final {order_i + 1}")
        FILLERS[order_i](prs.slides[slide_idx])

    # Delete slides NOT in KEEP_ORDER (from end to start)
    keep_set = set(KEEP_ORDER)
    for i in range(len(prs.slides) - 1, -1, -1):
        if i not in keep_set:
            delete_slide(prs, i)

    # After deletions, slides are in original relative order among survivors.
    # Survivors in ascending original index:
    survivors_asc = sorted(KEEP_ORDER)
    # We want KEEP_ORDER sequence. Map desired → current position
    current_positions = {orig: pos for pos, orig in enumerate(survivors_asc)}
    new_order = [current_positions[o] for o in KEEP_ORDER]
    reorder_slides(prs, new_order)

    assert len(prs.slides) == 10, len(prs.slides)

    # Write speaker notes markdown
    lines = ["# ملاحظات المتحدث — فهرس مناهج التفسير\n", "المجموع ≈ ٥ دقائق.\n"]
    for title, note in SPEAKER_NOTES:
        lines.append(f"## {title}\n\n{note}\n")
    NOTES_OUT.write_text("\n".join(lines), encoding="utf-8")

    print("Saving", PPTX_OUT)
    prs.save(str(PPTX_OUT))

    print("Compressing media…")
    compress_media_in_pptx(PPTX_OUT)

    # Re-open and gate-scan
    prs2 = Presentation(str(PPTX_OUT))
    print("Slide count:", len(prs2.slides))
    hits = scan_brackets(prs2)
    if hits:
        print("BRACKET HITS:")
        for h in hits:
            print(" ", h)
    else:
        print("Placeholder scan: 0")

    print("Per-slide min font sizes (≥12 pt gate):")
    tiny = scan_min_font_sizes(prs2, min_pt=12.0)
    if tiny:
        print("TINY FONT HITS:")
        for t in tiny:
            print(" ", t)
    else:
        print("Min font gate: OK (no run < 12 pt)")

    # Guide-slide heuristic: look for دليل الاستخدام
    for i, slide in enumerate(prs2.slides, 1):
        blob = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        if "دليل الاستخدام" in blob or "اقرأ الدليل" in blob:
            print(f"WARN: possible guide text on slide {i}")

    ok_pdf = export_pdf(PPTX_OUT, PDF_OUT)
    print("PDF export:", "OK" if ok_pdf else "FAILED")

    if PDF_OUT.exists():
        print("Checking PDF text overflow + tiny spans…")
        overflow = verify_pdf_overflow(PDF_OUT)
        if overflow:
            print("OVERFLOW / EDGE CLIP / TINY:")
            for o in overflow:
                print(" ", o)
        else:
            print("No text overflow vs slide edges (margin 8pt); no PDF span < 12 pt.")

    for p in (PPTX_OUT, PDF_OUT):
        if p.exists():
            print(f"SIZE {p.name}: {p.stat().st_size} ({p.stat().st_size/1024/1024:.2f} MB)")
        else:
            print(f"MISSING {p}")

    # Hard gates summary
    n_slides = len(prs2.slides)
    pptx_ok = PPTX_OUT.exists() and PPTX_OUT.stat().st_size < 10_000_000
    pdf_ok = PDF_OUT.exists() and PDF_OUT.stat().st_size < 10_000_000
    print(
        f"GATES: slides={n_slides}/10  pptx_lt_10mb={pptx_ok}  pdf_lt_10mb={pdf_ok}  "
        f"placeholders={len(hits)}  tiny_runs={len(tiny)}"
    )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
